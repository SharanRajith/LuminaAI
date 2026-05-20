import asyncio
import io
import json
import logging
import re
import os
import subprocess
import tempfile
import urllib.parse
from typing import Optional

import requests
from fastapi import FastAPI, File, HTTPException, UploadFile, Request, Header, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, field_validator
from dotenv import load_dotenv

from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from supabase import create_client, Client

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
logger = logging.getLogger("lumina")

# ─────────────────────────── Setup ────────────────────────────

SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
supabase: Optional[Client] = None
if SUPABASE_URL and SUPABASE_KEY:
    try:
        supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
        logger.info("Supabase client initialised.")
    except Exception as e:
        logger.warning("Supabase disabled — invalid credentials: %s", e)

def get_user_or_ip(request: Request) -> str:
    auth_header = request.headers.get("Authorization")
    if supabase and auth_header:
        try:
            token = auth_header.replace("Bearer ", "")
            user = supabase.auth.get_user(token)
            if user and user.user:
                return f"user:{user.user.id}"
        except Exception as e:
            logger.warning("Token verification failed in rate limiter: %s", e)
    return f"ip:{get_remote_address(request)}"

limiter = Limiter(key_func=get_user_or_ip)

app = FastAPI(title="Lumina AI API")
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

cors_origin = os.environ.get("CORS_ORIGIN", "*")
if cors_origin == "*":
    logger.warning(
        "CORS_ORIGIN is wildcard '*'. "
        "Set the CORS_ORIGIN environment variable to your frontend URL in production."
    )

app.add_middleware(
    CORSMiddleware,
    allow_origins=[cors_origin],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─────────────────────────── Models ───────────────────────────

_VALID_THEMES       = {"dark", "ocean", "sunset", "emerald", "corporate"}
_VALID_TONES        = {"professional", "casual", "academic", "inspirational", "humorous"}
_VALID_AUDIENCES    = {"general", "executive", "technical", "student", "expert"}
_VALID_LENGTHS      = {"short", "medium", "long"}
_VALID_REPORT_TYPES = {"business", "research", "technical", "academic", "marketing"}
_VALID_LANGUAGES    = {
    "english", "hindi", "spanish", "french", "german",
    "arabic", "portuguese", "chinese", "japanese", "korean",
}


class PresentationRequest(BaseModel):
    prompt: str = Field(..., min_length=1, max_length=5000)
    context: Optional[str] = Field(None, max_length=12000)
    model_name: Optional[str] = None
    theme: str = Field("dark")
    slide_count: int = Field(10, ge=3, le=50)
    audience: str = Field("general")
    tone: str = Field("professional")
    language: str = Field("english")

    @field_validator("theme")
    @classmethod
    def validate_theme(cls, v):
        if v not in _VALID_THEMES:
            raise ValueError(f"theme must be one of {sorted(_VALID_THEMES)}")
        return v

    @field_validator("tone")
    @classmethod
    def validate_tone(cls, v):
        if v not in _VALID_TONES:
            raise ValueError(f"tone must be one of {sorted(_VALID_TONES)}")
        return v

    @field_validator("audience")
    @classmethod
    def validate_audience(cls, v):
        if v not in _VALID_AUDIENCES:
            raise ValueError(f"audience must be one of {sorted(_VALID_AUDIENCES)}")
        return v

    @field_validator("language")
    @classmethod
    def validate_language(cls, v):
        if v not in _VALID_LANGUAGES:
            raise ValueError(f"language must be one of {sorted(_VALID_LANGUAGES)}")
        return v


class ReportRequest(BaseModel):
    prompt: str = Field(..., min_length=1, max_length=5000)
    context: Optional[str] = Field(None, max_length=12000)
    model_name: Optional[str] = None
    report_type: str = Field("business")
    tone: str = Field("professional")
    length: str = Field("medium")
    language: str = Field("english")

    @field_validator("report_type")
    @classmethod
    def validate_report_type(cls, v):
        if v not in _VALID_REPORT_TYPES:
            raise ValueError(f"report_type must be one of {sorted(_VALID_REPORT_TYPES)}")
        return v

    @field_validator("tone")
    @classmethod
    def validate_tone(cls, v):
        if v not in _VALID_TONES:
            raise ValueError(f"tone must be one of {sorted(_VALID_TONES)}")
        return v

    @field_validator("length")
    @classmethod
    def validate_length(cls, v):
        if v not in _VALID_LENGTHS:
            raise ValueError(f"length must be one of {sorted(_VALID_LENGTHS)}")
        return v

    @field_validator("language")
    @classmethod
    def validate_language(cls, v):
        if v not in _VALID_LANGUAGES:
            raise ValueError(f"language must be one of {sorted(_VALID_LANGUAGES)}")
        return v


class SlideRegenerateRequest(BaseModel):
    presentation_title: str = Field(..., min_length=1, max_length=200)
    slide_type: str = Field("content")
    topic: str = Field(..., min_length=3, max_length=500)
    tone: str = Field("professional")
    audience: str = Field("general")

    @field_validator("tone")
    @classmethod
    def validate_tone(cls, v):
        if v not in _VALID_TONES:
            raise ValueError(f"tone must be one of {sorted(_VALID_TONES)}")
        return v

    @field_validator("audience")
    @classmethod
    def validate_audience(cls, v):
        if v not in _VALID_AUDIENCES:
            raise ValueError(f"audience must be one of {sorted(_VALID_AUDIENCES)}")
        return v


class ExportRequest(BaseModel):
    presentation_data: dict

# ─────────────────────────── Helpers ──────────────────────────

def extract_json(text: str) -> dict:
    """Extract the first complete JSON object from LLM response text.

    Uses brace-matching instead of a greedy regex so nested objects
    and multiple JSON blocks are handled correctly.
    """
    def clean_obj(obj):
        if isinstance(obj, dict):
            return {k: clean_obj(v) for k, v in obj.items()}
        elif isinstance(obj, list):
            return [clean_obj(v) for v in obj]
        elif isinstance(obj, str):
            return re.sub(r'<[^>]*>', '', obj)
        return obj

    start = text.find('{')
    if start == -1:
        raise ValueError("No JSON object found in AI response")

    depth = 0
    in_string = False
    escape_next = False

    for i, ch in enumerate(text[start:], start):
        if escape_next:
            escape_next = False
            continue
        if ch == '\\' and in_string:
            escape_next = True
            continue
        if ch == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if ch == '{':
            depth += 1
        elif ch == '}':
            depth -= 1
            if depth == 0:
                candidate = text[start:i + 1]
                try:
                    data = json.loads(candidate)
                    return clean_obj(data)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"Malformed JSON in AI response: {exc}") from exc

    raise ValueError("No complete JSON object found in AI response")


def smart_truncate(text: str, max_chars: int = 6000) -> str:
    """Truncate text at a paragraph or sentence boundary instead of mid-word."""
    if len(text) <= max_chars:
        return text
    window = text[:max_chars]
    last_para = window.rfind('\n\n')
    if last_para > max_chars * 0.8:
        return window[:last_para].strip()
    for punct in ('. ', '! ', '? '):
        last_sent = window.rfind(punct)
        if last_sent > max_chars * 0.8:
            return window[:last_sent + 1].strip()
    return window.rstrip()


def generate_with_groq(prompt: str, model_name: Optional[str] = None):
    api_key = os.environ.get("GROQ_API_KEY", "")
    if not api_key:
        raise ValueError("GROQ_API_KEY is not set. Get a free key at https://console.groq.com")
    model = model_name or "llama-3.3-70b-versatile"
    response = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json={
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "response_format": {"type": "json_object"},
            "temperature": 0.7,
        },
        timeout=60,
    )
    if response.status_code != 200:
        raise ValueError(f"Groq API error {response.status_code}: {response.text}")
    return response.json()["choices"][0]["message"]["content"]


def verify_token(authorization: Optional[str] = Header(None)) -> Optional[str]:
    """Optional JWT verification. Returns user_id if valid."""
    if not supabase or not authorization:
        return None
    try:
        token = authorization.replace("Bearer ", "")
        user = supabase.auth.get_user(token)
        if user and user.user:
            _ensure_admin_premium(user.user.id, user.user.email or "")
            return user.user.id
        return None
    except Exception as e:
        logger.warning("Token verification failed: %s", e)
        return None

# ─────────────────────────── Tier System ──────────────────────

FREE_TIER_LIMIT  = 5
ADMIN_SECRET_KEY = os.environ.get("ADMIN_SECRET_KEY", "")
_ADMIN_EMAILS_RAW = os.environ.get("ADMIN_EMAILS", "sharanrajithk@gmail.com,madhurahegde475@gmail.com")
ADMIN_EMAILS      = {e.strip().lower() for e in _ADMIN_EMAILS_RAW.split(",") if e.strip()}

def _ensure_admin_premium(user_id: str, email: str):
    """If this user is the admin, silently upgrade their profile to premium."""
    if email.lower() not in ADMIN_EMAILS:
        return
    try:
        res = supabase.table("profiles").select("tier").eq("id", user_id).execute()
        if res.data and res.data[0]["tier"] == "premium":
            return  # already premium, nothing to do
        supabase.table("profiles").upsert({
            "id": user_id, "tier": "premium", "generations_used": 0
        }).execute()
        logger.info("Auto-granted premium to admin account %s", email)
    except Exception as e:
        logger.warning("Failed to auto-grant admin premium: %s", e)

def _get_or_create_profile(user_id: str) -> dict:
    res = supabase.table("profiles").select("*").eq("id", user_id).execute()
    if res.data:
        return res.data[0]
    new = supabase.table("profiles").insert({
        "id": user_id, "tier": "free", "generations_used": 0
    }).execute()
    return new.data[0]

def check_tier_and_increment(user_id: Optional[str]):
    """Raise 403 if free tier limit reached; otherwise increment usage counter."""
    if not supabase or not user_id:
        return  # unauthenticated users pass through (IP rate-limiter still applies)
    profile = _get_or_create_profile(user_id)
    if profile["tier"] == "premium":
        return
    used = profile["generations_used"]
    if used >= FREE_TIER_LIMIT:
        raise HTTPException(
            status_code=403,
            detail=f"limit_reached:{used}:{FREE_TIER_LIMIT}"
        )
    supabase.table("profiles").update({"generations_used": used + 1}).eq("id", user_id).execute()

# ─────────────────────────── Routes ───────────────────────────

@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/user/profile")
async def get_user_profile(user_id: Optional[str] = Depends(verify_token)):
    if not user_id:
        raise HTTPException(status_code=401, detail="Not authenticated")
    profile = _get_or_create_profile(user_id)
    return {
        "tier":             profile["tier"],
        "generations_used": profile["generations_used"],
        "limit":            FREE_TIER_LIMIT if profile["tier"] == "free" else None,
    }


@app.post("/admin/grant-premium")
async def grant_premium(request: Request):
    key = request.headers.get("X-Admin-Key", "")
    if not ADMIN_SECRET_KEY or key != ADMIN_SECRET_KEY:
        raise HTTPException(status_code=403, detail="Unauthorized")
    body = await request.json()
    user_id = body.get("user_id", "").strip()
    if not user_id:
        raise HTTPException(status_code=400, detail="user_id required")
    supabase.table("profiles").upsert({"id": user_id, "tier": "premium"}).execute()
    logger.info("Granted premium to user %s", user_id)
    return {"status": "ok", "user_id": user_id, "tier": "premium"}


def verify_admin(authorization: Optional[str] = Header(None)) -> str:
    """Dependency: verifies the caller is the admin (by email match)."""
    if not supabase or not authorization:
        raise HTTPException(status_code=403, detail="Unauthorized")
    try:
        token = authorization.replace("Bearer ", "")
        user  = supabase.auth.get_user(token)
        if user and user.user and (user.user.email or "").lower() in ADMIN_EMAILS:
            return user.user.id
    except Exception:
        pass
    raise HTTPException(status_code=403, detail="Unauthorized")


_IMAGE_EXTS    = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
_IMAGE_MIMES   = {".jpg": "image/jpeg", ".jpeg": "image/jpeg",
                  ".png": "image/png",  ".webp": "image/webp", ".gif": "image/gif"}
_DOC_EXTS      = {".pdf", ".docx", ".txt"}

@app.post("/upload/document")
async def upload_document(file: UploadFile = File(...), _user_id: Optional[str] = Depends(verify_token)):
    filename = (file.filename or "").lower()
    ext = next((e for e in _IMAGE_EXTS | _DOC_EXTS if filename.endswith(e)), None)
    if not ext:
        raise HTTPException(status_code=400,
            detail="Unsupported file type. Upload PDF, DOCX, TXT, JPG, PNG, or WEBP.")
    content = await file.read()
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large (max 15 MB).")
    text = ""
    try:
        if ext in _IMAGE_EXTS:
            # Use Groq Vision to describe the image
            api_key = os.environ.get("GROQ_API_KEY", "")
            if not api_key:
                raise HTTPException(status_code=500, detail="GROQ_API_KEY not set.")
            import base64
            mime      = _IMAGE_MIMES.get(ext, "image/jpeg")
            b64_image = base64.b64encode(content).decode("utf-8")
            vision_res = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json={
                    "model": "llama-3.2-11b-vision-preview",
                    "messages": [{
                        "role": "user",
                        "content": [
                            {"type": "image_url",
                             "image_url": {"url": f"data:{mime};base64,{b64_image}"}},
                            {"type": "text",
                             "text": (
                                "Analyze this image thoroughly. Extract all visible text, numbers, "
                                "data, charts, and key information. Describe what you see in detail — "
                                "including subjects, context, any labels or captions. "
                                "This description will be used to create a presentation or report."
                             )}
                        ]
                    }],
                    "max_tokens": 1500,
                },
                timeout=30,
            )
            if vision_res.status_code != 200:
                raise HTTPException(status_code=422,
                    detail=f"Image analysis failed: {vision_res.text[:200]}")
            text = vision_res.json()["choices"][0]["message"]["content"].strip()
        elif ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    text += (page.extract_text() or "") + "\n"
        elif ext == ".docx":
            from docx import Document as DocxDoc
            doc = DocxDoc(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            text = content.decode("utf-8", errors="ignore")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not parse file: {e}")
    text = text.strip()[:10000]
    if not text:
        raise HTTPException(status_code=422, detail="No content could be extracted from this file.")
    is_image = ext in _IMAGE_EXTS
    return {"text": text, "filename": file.filename, "chars": len(text), "is_image": is_image}


@app.get("/admin/stats")
async def admin_stats(_admin_id: str = Depends(verify_admin)):
    profiles_res  = supabase.table("profiles").select("tier,generations_used").execute()
    profiles      = profiles_res.data or []
    creations_res = supabase.table("creations").select("type").execute()
    creations     = creations_res.data or []
    return {
        "total_users":        len(profiles),
        "premium_users":      sum(1 for p in profiles if p["tier"] == "premium"),
        "free_users":         sum(1 for p in profiles if p["tier"] == "free"),
        "total_generations":  sum(p["generations_used"] for p in profiles),
        "presentations":      sum(1 for c in creations if c["type"] == "presentation"),
        "reports":            sum(1 for c in creations if c["type"] == "report"),
    }


@app.get("/admin/users")
async def admin_users(_admin_id: str = Depends(verify_admin)):
    profiles_res = supabase.table("profiles").select("*").order("created_at", desc=True).execute()
    profiles     = {p["id"]: p for p in (profiles_res.data or [])}
    try:
        auth_list  = supabase.auth.admin.list_users()
        auth_users = getattr(auth_list, "users", auth_list) or []
    except Exception:
        auth_users = []
    users = []
    for au in auth_users:
        p = profiles.get(au.id, {})
        users.append({
            "id":               au.id,
            "email":            au.email or "—",
            "tier":             p.get("tier", "free"),
            "generations_used": p.get("generations_used", 0),
            "created_at":       p.get("created_at") or str(getattr(au, "created_at", "")),
        })
    return {"users": users}


@app.post("/admin/set-tier")
async def admin_set_tier(body: dict, _admin_id: str = Depends(verify_admin)):
    user_id = body.get("user_id", "").strip()
    tier    = body.get("tier", "").strip()
    if not user_id or tier not in ("free", "premium"):
        raise HTTPException(status_code=400, detail="user_id and tier ('free'|'premium') required")
    supabase.table("profiles").upsert({"id": user_id, "tier": tier}).execute()
    logger.info("Admin set tier=%s for user %s", tier, user_id)
    return {"status": "ok", "user_id": user_id, "tier": tier}


@app.post("/generate/presentation")
@limiter.limit("3/day")
async def generate_presentation(req: PresentationRequest, request: Request, user_id: Optional[str] = Depends(verify_token)):
    check_tier_and_increment(user_id)
    try:
        lang_instruction = f"- Language: Write ALL content (titles, bullets, notes, subtitles) in {req.language.capitalize()}. Do NOT mix languages.\n" if req.language != "english" else ""
        if req.context:
            content_section = f"DOCUMENT CONTENT (base the presentation on this):\n{req.context[:8000]}\n\nUSER INSTRUCTIONS: {req.prompt}"
        else:
            content_section = f"TOPIC: {req.prompt}"
        prompt = f"""You are a senior consultant creating an authoritative, information-dense presentation. Think McKinsey-quality slide deck — every slide must deliver real, specific, actionable knowledge.

REQUIREMENTS:
- Exactly {req.slide_count} slides
- Audience: {req.audience}
- Tone: {req.tone}
{lang_instruction}
CONTENT RULES (strictly follow):
- Each content slide: 3-5 bullets. Every bullet must be 12-25 words with SPECIFIC facts, data, mechanisms, or named examples. No vague one-liners.
- Include real numbers, percentages, timelines, statistics, or case examples wherever relevant.
- NO quote slides — every slide must contain original analysis or factual information.
- Stats slides: use realistic, specific numerical data with descriptive labels and context.
- Conclusion: concrete, actionable takeaways — not vague summaries.
- Agenda slide: list the actual section names of the presentation.

IMAGE POSITIONING:
- "full"  → ONLY for the title slide (full-bleed background image)
- "right" → most content slides (image panel right, text left)
- "left"  → use 1-2 times for visual variety
- "none"  → agenda slide and stats slide (data looks cluttered with image)

{content_section}

Return ONLY valid JSON. No markdown fences. No HTML or Markdown inside string values. Plain text only.

{{
  "title": "...",
  "subtitle": "...",
  "theme": "{req.theme}",
  "slides": [
    {{ "type": "title",     "title": "...", "subtitle": "...", "image_position": "full",  "notes": "Speaker intro note..." }},
    {{ "type": "agenda",    "title": "Agenda", "items": ["Section 1", "Section 2", "Section 3"], "image_position": "none",  "notes": "..." }},
    {{ "type": "content",   "title": "...", "bullets": ["Specific fact with detail and context, about 15 words here", "Another substantive point with real data or named example"], "image_position": "right", "notes": "..." }},
    {{ "type": "stats",     "title": "...", "stats": [{{"value":"42%","label":"Adoption Rate","description":"Among Fortune 500 companies in 2024"}},{{"value":"3.2×","label":"ROI","description":"Average return within 18 months"}}], "image_position": "none",  "notes": "..." }},
    {{ "type": "content",   "title": "...", "bullets": ["...", "..."], "image_position": "left",  "notes": "..." }},
    {{ "type": "conclusion","title": "Key Takeaways", "bullets": ["Actionable takeaway with specific next step", "..."], "image_position": "right", "notes": "..." }}
  ]
}}"""

        response_text = generate_with_groq(prompt, req.model_name)
        data = extract_json(response_text)

        if supabase and user_id:
            try:
                supabase.table("creations").insert({
                    "user_id": user_id,
                    "type": "presentation",
                    "prompt": req.prompt,
                    "title": data.get("title", "Untitled"),
                    "output": data,
                }).execute()
            except Exception as e:
                logger.error("Failed to save presentation for user %s: %s", user_id, e)

        return data

    except Exception as e:
        logger.exception("Presentation generation failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/generate/report")
@limiter.limit("3/day")
async def generate_report(req: ReportRequest, request: Request, user_id: Optional[str] = Depends(verify_token)):
    check_tier_and_increment(user_id)
    try:
        section_count = {"short": "4-5", "medium": "6-8", "long": "9-12"}.get(req.length, "6-8")
        lang_instruction = f"- Language: Write ALL content in {req.language.capitalize()}. Do NOT mix languages.\n" if req.language != "english" else ""
        if req.context:
            content_section = f"DOCUMENT CONTENT (base the report on this):\n{req.context[:8000]}\n\nUSER INSTRUCTIONS: {req.prompt}"
        else:
            content_section = f"TOPIC: {req.prompt}"

        prompt = f"""You are an expert analyst and report writer.
Write a comprehensive, authoritative {req.report_type} report on the topic below.

REQUIREMENTS:
- Tone: {req.tone}
- Length: {section_count} sections
{lang_instruction}- Include executive summary, body sections, conclusion, and recommendations

{content_section}

Return ONLY valid JSON (no markdown fences) matching this schema exactly.
CRITICAL: Do NOT use HTML tags (like <h1>) or Markdown inside JSON values. Use plain text only!
{{
  "title": "...",
  "subtitle": "...",
  "report_type": "{req.report_type}",
  "date": "May 2025",
  "executive_summary": "3-4 sentence summary of the entire report...",
  "key_highlights": [
    {{"label": "...", "value": "..."}}
  ],
  "sections": [
    {{
      "heading": "...",
      "content": "Multiple detailed paragraphs of substantive content...",
      "key_points": ["...", "..."],
      "subsections": [
        {{"heading": "...", "content": "..."}}
      ]
    }}
  ],
  "conclusion": "Strong concluding paragraph...",
  "recommendations": ["Actionable recommendation 1", "...", "..."]
}}"""

        response_text = generate_with_groq(prompt, req.model_name)
        data = extract_json(response_text)

        if supabase and user_id:
            try:
                supabase.table("creations").insert({
                    "user_id": user_id,
                    "type": "report",
                    "prompt": req.prompt,
                    "title": data.get("title", "Untitled"),
                    "output": data,
                }).execute()
            except Exception as e:
                logger.error("Failed to save report for user %s: %s", user_id, e)

        return data

    except Exception as e:
        logger.exception("Report generation failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/generate/slide")
@limiter.limit("20/day")
async def regenerate_slide(req: SlideRegenerateRequest, request: Request):
    """Regenerate a single slide — used by the inline editor."""
    try:
        prompt = f"""You are an expert presentation designer.
Regenerate ONE slide for a presentation titled "{req.presentation_title}".

Slide topic: {req.topic}
Slide type: {req.slide_type}
Audience: {req.audience}
Tone: {req.tone}

Return ONLY a single JSON slide object (no array, no wrapper). Use plain text only, no HTML or Markdown.
Choose image_position from: "full" | "right" | "left" | "none"

For type "content":  {{"type":"content","title":"...","bullets":["...","...","..."],"image_position":"right","notes":"..."}}
For type "quote":    {{"type":"quote","quote":"...","author":"...","image_position":"full","notes":"..."}}
For type "stats":    {{"type":"stats","title":"...","stats":[{{"value":"...","label":"...","description":"..."}}],"image_position":"full","notes":"..."}}
For type "title":    {{"type":"title","title":"...","subtitle":"...","image_position":"full","notes":"..."}}

Make it substantive and engaging."""

        response_text = generate_with_groq(prompt, None)
        slide = extract_json(response_text)
        return slide
    except Exception as e:
        logger.exception("Single slide regeneration failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/pptx")
async def export_pptx(req: ExportRequest):
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        data = req.presentation_data
        theme = data.get("theme", "dark")

        THEMES = {
            "dark":      {"bg": (0x0F,0x0F,0x23), "bg2": (0x1A,0x1A,0x35), "acc": (0x6C,0x63,0xFF), "acc2": (0xFF,0x65,0x84), "txt": (0xF0,0xF0,0xFF), "txt2": (0xA0,0xA0,0xC0)},
            "ocean":     {"bg": (0x07,0x14,0x28), "bg2": (0x0D,0x20,0x40), "acc": (0x00,0xC6,0xFB), "acc2": (0x00,0x5B,0xEA), "txt": (0xE0,0xF4,0xFF), "txt2": (0xC0,0xE8,0xFF)},
            "sunset":    {"bg": (0x1A,0x08,0x13), "bg2": (0x2D,0x10,0x20), "acc": (0xFF,0x65,0x84), "acc2": (0xFF,0x9F,0x43), "txt": (0xFF,0xF0,0xF3), "txt2": (0xFF,0xD0,0xD8)},
            "emerald":   {"bg": (0x02,0x14,0x08), "bg2": (0x04,0x1F,0x0F), "acc": (0x00,0xC9,0xA7), "acc2": (0x00,0xE6,0x76), "txt": (0xE0,0xFF,0xF2), "txt2": (0xA0,0xFF,0xD0)},
            "corporate": {"bg": (0xF4,0xF6,0xFB), "bg2": (0xE8,0xEC,0xF5), "acc": (0x25,0x63,0xEB), "acc2": (0x7C,0x3A,0xED), "txt": (0x11,0x18,0x27), "txt2": (0x4B,0x55,0x63)},
        }
        TC  = THEMES.get(theme, THEMES["dark"])
        BG  = RGBColor(*TC["bg"])
        BG2 = RGBColor(*TC["bg2"])
        ACC = RGBColor(*TC["acc"])
        ACC2= RGBColor(*TC["acc2"])
        TXT = RGBColor(*TC["txt"])
        TXT2= RGBColor(*TC["txt2"])

        prs = Presentation()
        prs.slide_width  = Inches(13.33)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]

        def add_rect(slide, x, y, w, h, color):
            s = slide.shapes.add_shape(1, x, y, w, h)
            s.fill.solid()
            s.fill.fore_color.rgb = color
            s.line.fill.background()
            return s

        def add_text(slide, text, x, y, w, h, size, bold=False, color=None, italic=False, align=PP_ALIGN.LEFT):
            if not color: color = TXT
            box = slide.shapes.add_textbox(x, y, w, h)
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = str(text)
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.italic = italic
            p.font.color.rgb = color
            p.alignment = align

        def set_bg(slide):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = BG
            add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.08), ACC)
            add_rect(slide, Inches(0), Inches(7.3), Inches(13.33), Inches(0.2), BG2)

        async def fetch_image(sd: dict):
            """Fetch a Pollinations AI image in a thread so the event loop isn't blocked."""
            title   = sd.get("title", "")
            bullets = sd.get("bullets", sd.get("items", []))
            extra   = (bullets[0] if bullets else "")[:70]
            desc    = f"{title}, {extra}" if extra else title
            encoded_prompt = urllib.parse.quote(
                f"wide angle landscape scene, {desc}, professional stock photography, "
                f"cinematic lighting, high resolution, vivid colors, no text, no watermark, no portrait"
            )
            url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=1280&height=720&nologo=true"

            def _fetch():
                try:
                    r = requests.get(url, timeout=12)
                    if r.status_code == 200:
                        return io.BytesIO(r.content)
                except Exception as e:
                    logger.warning("Image fetch failed for '%s': %s", title, e)
                return None

            return await asyncio.to_thread(_fetch)

        slides = data.get("slides", [])
        total  = len(slides)

        # Fetch all images concurrently instead of one-by-one
        slide_images = await asyncio.gather(
            *[fetch_image(sd) for sd in slides]
        )

        for idx, sd in enumerate(slides):
            slide = prs.slides.add_slide(blank)
            set_bg(slide)
            stype     = sd.get("type", "content")
            title_txt = sd.get("title", "")
            img_data  = slide_images[idx]

            add_text(slide, f"{idx+1} / {total}", Inches(0.4), Inches(7.1), Inches(1.4), Inches(0.3),
                     10, color=TXT2, align=PP_ALIGN.LEFT)
            add_text(slide, "Made with Lumina AI", Inches(9.5), Inches(7.1), Inches(3.5), Inches(0.3),
                     9, color=TXT2, align=PP_ALIGN.RIGHT)

            if stype == "title":
                if img_data:
                    try:
                        slide.shapes.add_picture(img_data, Inches(7.0), Inches(0), Inches(6.33), Inches(7.5))
                        add_rect(slide, Inches(6.9), Inches(0), Inches(0.2), Inches(7.5), ACC)
                    except Exception:
                        pass

                add_text(slide, title_txt, Inches(0.7), Inches(2.0), Inches(6.0), Inches(2.6),
                         48, bold=True, color=TXT)
                sub = sd.get("subtitle", "")
                if sub:
                    add_text(slide, sub, Inches(0.7), Inches(4.7), Inches(6.0), Inches(1.4),
                             22, color=TXT2)

            elif stype == "quote":
                if img_data:
                    try:
                        slide.shapes.add_picture(img_data, Inches(7.5), Inches(1.0), Inches(5.0), Inches(5.5))
                    except Exception:
                        pass

                add_text(slide, "“", Inches(0.6), Inches(1.0), Inches(2), Inches(1.5),
                         100, bold=True, color=RGBColor(0x35,0x35,0x65))
                add_text(slide, sd.get("quote",""), Inches(0.8), Inches(2.2), Inches(6.2), Inches(3.5),
                         26, italic=True, color=TXT)
                add_rect(slide, Inches(0.8), Inches(5.7), Inches(1.5), Inches(0.05), ACC)
                add_text(slide, f"— {sd.get('author','')}", Inches(0.8), Inches(5.85), Inches(6.0), Inches(0.7),
                         18, bold=True, color=ACC)

            elif stype == "stats":
                if img_data:
                    try:
                        slide.shapes.add_picture(img_data, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
                        add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), RGBColor(0,0,0))
                        slide.shapes[-1].fill.transparency = 0.75
                    except Exception:
                        pass

                add_text(slide, title_txt, Inches(0.5), Inches(0.5), Inches(12), Inches(0.85),
                         38, bold=True, color=TXT, align=PP_ALIGN.CENTER)
                add_rect(slide, Inches(5.66), Inches(1.4), Inches(2.0), Inches(0.06), ACC)

                stats = sd.get("stats", [])
                n     = max(len(stats), 1)
                col_w = 12.2 / n
                for i, st in enumerate(stats[:4]):
                    x = Inches(0.55 + i * col_w)
                    w = Inches(col_w - 0.15)
                    add_text(slide, st.get("value",""), x, Inches(2.8), w, Inches(1.5),
                             54, bold=True, color=ACC, align=PP_ALIGN.CENTER)
                    add_text(slide, st.get("label",""), x, Inches(4.4), w, Inches(0.65),
                             20, bold=True, color=TXT, align=PP_ALIGN.CENTER)

            else:  # content / agenda / conclusion / image
                bullets = sd.get("bullets", sd.get("items", []))
                if not bullets and sd.get("caption"):
                    bullets = [sd.get("caption")]

                layout_style = idx % 3

                if layout_style == 0:
                    add_rect(slide, Inches(0), Inches(0.08), Inches(0.12), Inches(7.22), ACC2)
                    add_text(slide, title_txt, Inches(0.4), Inches(0.2), Inches(6.5), Inches(0.9), 32, bold=True, color=TXT)
                    add_rect(slide, Inches(0.4), Inches(1.2), Inches(2.0), Inches(0.06), ACC)
                    if img_data:
                        try: slide.shapes.add_picture(img_data, Inches(7.5), Inches(1.0), Inches(5.33), Inches(5.5))
                        except Exception: pass
                    cx, cw = Inches(0.5), Inches(6.5)

                elif layout_style == 1:
                    add_rect(slide, Inches(13.21), Inches(0.08), Inches(0.12), Inches(7.22), ACC2)
                    add_text(slide, title_txt, Inches(6.5), Inches(0.2), Inches(6.5), Inches(0.9), 32, bold=True, color=TXT)
                    add_rect(slide, Inches(6.5), Inches(1.2), Inches(2.0), Inches(0.06), ACC)
                    if img_data:
                        try: slide.shapes.add_picture(img_data, Inches(0.5), Inches(1.0), Inches(5.33), Inches(5.5))
                        except Exception: pass
                    cx, cw = Inches(6.5), Inches(6.5)

                else:
                    if img_data:
                        try:
                            slide.shapes.add_picture(img_data, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
                            add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), RGBColor(0,0,0))
                            slide.shapes[-1].fill.transparency = 0.82
                        except Exception:
                            pass
                    add_rect(slide, Inches(0), Inches(0.08), Inches(0.12), Inches(7.22), ACC2)
                    add_text(slide, title_txt, Inches(0.4), Inches(0.2), Inches(12.0), Inches(0.9), 36, bold=True, color=TXT)
                    add_rect(slide, Inches(0.4), Inches(1.2), Inches(2.0), Inches(0.06), ACC)
                    cx, cw = Inches(0.5), Inches(12.0)

                for i, b in enumerate(bullets):
                    y = Inches(1.6 + i * 0.95)
                    if y > Inches(6.5): break
                    add_rect(slide, cx, y + Inches(0.15), Inches(0.1), Inches(0.1), ACC)
                    add_text(slide, b, cx + Inches(0.22), y, cw - Inches(0.22), Inches(0.9), 16, color=TXT)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": 'attachment; filename="lumina_presentation.pptx"'},
        )

    except Exception as e:
        logger.exception("PPTX export failed")
        raise HTTPException(status_code=500, detail=str(e))


def build_latex_source(data: dict) -> str:
    """Convert report JSON dict into a LaTeX document string."""
    def tex(s: str) -> str:
        for ch, rep in [("&","\\&"),("%","\\%"),("$","\\$"),("#","\\#"),
                        ("_","\\_"),("{","\\{"),("}","\\}"),
                        ("~","\\textasciitilde{}"),("^","\\textasciicircum{}")]:
            s = s.replace(ch, rep)
        return s

    title    = tex(data.get("title", "Report"))
    subtitle = tex(data.get("subtitle", ""))
    date     = tex(data.get("date", ""))
    sections = data.get("sections", [])

    lines = [
        r"\documentclass[12pt,a4paper]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{lmodern}",
        r"\usepackage[margin=2.5cm]{geometry}",
        r"\usepackage[dvipsnames]{xcolor}",
        r"\usepackage[colorlinks=true,linkcolor=accent,urlcolor=accent]{hyperref}",
        r"\usepackage{fancyhdr}",
        r"\usepackage{parskip}",
        r"\definecolor{accent}{RGB}{108,99,255}",
        r"\makeatletter",
        r"\renewcommand\section{\@startsection{section}{1}{\z@}{-3.5ex \@plus -1ex \@minus -.2ex}{2.3ex \@plus.2ex}{\Large\bfseries\color{accent}}}",
        r"\makeatother",
        r"\pagestyle{fancy}",
        rf"\fancyhead[L]{{\textbf{{{title}}}}}",
        rf"\fancyhead[R]{{{date}}}",
        r"\fancyfoot[C]{\thepage}",
        r"\fancyfoot[R]{\small\textcolor{accent}{Made with Lumina AI}}",
        r"\begin{document}",
        r"\begin{titlepage}\centering",
        r"\vspace*{3cm}",
        rf"{{\Huge\bfseries\color{{accent}} {title}}}\\[1cm]",
        rf"{{\Large\itshape {subtitle}}}\\[0.8cm]",
        rf"{{\large {date}}}",
        r"\vfill{\itshape Generated by Lumina AI}",
        r"\end{titlepage}",
        r"\tableofcontents\newpage",
    ]

    if data.get("executive_summary"):
        lines += [r"\section*{Executive Summary}",
                  r"\addcontentsline{toc}{section}{Executive Summary}",
                  tex(data["executive_summary"]), ""]

    if data.get("key_highlights"):
        lines += [r"\section*{Key Highlights}",
                  r"\addcontentsline{toc}{section}{Key Highlights}",
                  r"\begin{itemize}"]
        for h in data["key_highlights"]:
            lines.append(rf"\item \textbf{{{tex(h.get('label',''))}}}: {tex(h.get('value',''))}")
        lines += [r"\end{itemize}", ""]

    for s in sections:
        lines.append(rf"\section{{{tex(s.get('heading',''))}}}")
        if s.get("content"):
            lines.append(tex(s["content"]))
        if s.get("key_points"):
            lines += [r"\begin{itemize}"]
            for p in s["key_points"]:
                lines.append(rf"\item {tex(p)}")
            lines += [r"\end{itemize}"]
        for sub in s.get("subsections", []):
            lines.append(rf"\subsection{{{tex(sub.get('heading',''))}}}")
            lines.append(tex(sub.get("content", "")))
        lines.append("")

    if data.get("conclusion"):
        lines += [r"\section*{Conclusion}",
                  r"\addcontentsline{toc}{section}{Conclusion}",
                  tex(data["conclusion"]), ""]

    if data.get("recommendations"):
        lines += [r"\section*{Recommendations}",
                  r"\addcontentsline{toc}{section}{Recommendations}",
                  r"\begin{enumerate}"]
        for r_text in data["recommendations"]:
            lines.append(rf"\item {tex(r_text)}")
        lines += [r"\end{enumerate}"]

    lines.append(r"\end{document}")
    return "\n".join(lines)


def build_docx(data: dict) -> io.BytesIO:
    """Convert report JSON dict into a Word (.docx) document."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    ACCENT = RGBColor(0x6C, 0x63, 0xFF)

    doc = Document()

    # Page margins + footer
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.2)
        section.right_margin  = Inches(1.2)
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        footer_para = section.footer.paragraphs[0]
        footer_para.clear()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr = footer_para.add_run("Made with Lumina AI")
        fr.font.size = Pt(8)
        fr.font.color.rgb = ACCENT
        fr.font.bold = True

    # Title page
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(data.get("title", "Report"))
    tr.bold = True
    tr.font.size = Pt(28)
    tr.font.color.rgb = ACCENT

    if data.get("subtitle"):
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(data["subtitle"])
        sr.italic = True
        sr.font.size = Pt(15)

    if data.get("date"):
        dp = doc.add_paragraph()
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dp.add_run(data["date"]).font.size = Pt(12)

    doc.add_paragraph()
    gp = doc.add_paragraph()
    gp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gr = gp.add_run("Generated by Lumina AI")
    gr.italic = True
    gr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_page_break()

    if data.get("executive_summary"):
        h = doc.add_heading("Executive Summary", level=1)
        h.runs[0].font.color.rgb = ACCENT
        doc.add_paragraph(data["executive_summary"])

    if data.get("key_highlights"):
        h = doc.add_heading("Key Highlights", level=1)
        h.runs[0].font.color.rgb = ACCENT
        for hl in data["key_highlights"]:
            p = doc.add_paragraph(style="List Bullet")
            r1 = p.add_run(hl.get("label", "") + ": ")
            r1.bold = True
            p.add_run(hl.get("value", ""))

    for sec in data.get("sections", []):
        h = doc.add_heading(sec.get("heading", ""), level=1)
        if h.runs:
            h.runs[0].font.color.rgb = ACCENT
        if sec.get("content"):
            for para_text in sec["content"].split("\n\n"):
                if para_text.strip():
                    doc.add_paragraph(para_text.strip())
        if sec.get("key_points"):
            for pt in sec["key_points"]:
                doc.add_paragraph(pt, style="List Bullet")
        for sub in sec.get("subsections", []):
            doc.add_heading(sub.get("heading", ""), level=2)
            if sub.get("content"):
                doc.add_paragraph(sub["content"])

    if data.get("conclusion"):
        h = doc.add_heading("Conclusion", level=1)
        if h.runs:
            h.runs[0].font.color.rgb = ACCENT
        doc.add_paragraph(data["conclusion"])

    if data.get("recommendations"):
        h = doc.add_heading("Recommendations", level=1)
        if h.runs:
            h.runs[0].font.color.rgb = ACCENT
        for rec in data["recommendations"]:
            doc.add_paragraph(rec, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@app.post("/export/report-docx")
async def export_report_docx(req: ExportRequest):
    """Generate a Word (.docx) document from report JSON."""
    try:
        buf  = build_docx(req.presentation_data)
        name = req.presentation_data.get("title", "report").replace(" ", "_")
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{name}.docx"'},
        )
    except Exception as e:
        logger.exception("DOCX export failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/report-latex")
async def export_report_latex(req: ExportRequest):
    """Return the raw LaTeX source (.tex) for the report."""
    try:
        src = build_latex_source(req.presentation_data).encode("utf-8")
        name = req.presentation_data.get("title", "report").replace(" ", "_")
        return StreamingResponse(
            io.BytesIO(src),
            media_type="application/x-tex",
            headers={"Content-Disposition": f'attachment; filename="{name}.tex"'},
        )
    except Exception as e:
        logger.exception("LaTeX source export failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/report-pdf")
async def export_report_pdf(req: ExportRequest):
    """Compile LaTeX → PDF with pdflatex and return the PDF."""
    try:
        tex_src = build_latex_source(req.presentation_data)
        name    = req.presentation_data.get("title", "report").replace(" ", "_")

        with tempfile.TemporaryDirectory() as tmpdir:
            tex_path = os.path.join(tmpdir, "report.tex")
            pdf_path = os.path.join(tmpdir, "report.pdf")

            with open(tex_path, "w", encoding="utf-8") as f:
                f.write(tex_src)

            cmd = ["pdflatex", "-interaction=nonstopmode",
                   "-output-directory", tmpdir, tex_path]

            # Run twice so TOC page numbers resolve correctly
            for _ in range(2):
                result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)

            if not os.path.exists(pdf_path):
                log_tail = result.stdout[-2000:] if result.stdout else result.stderr[-2000:]
                raise ValueError(f"pdflatex failed:\n{log_tail}")

            with open(pdf_path, "rb") as f:
                pdf_bytes = f.read()

        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'},
        )
    except Exception as e:
        logger.exception("LaTeX PDF compilation failed")
        raise HTTPException(status_code=500, detail=str(e))



from fastapi.staticfiles import StaticFiles
frontend_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "frontend"))
if os.path.exists(frontend_path):
    app.mount("/", StaticFiles(directory=frontend_path, html=True), name="frontend")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
